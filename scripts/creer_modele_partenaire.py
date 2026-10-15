"""Crée le modèle de fichier à envoyer aux partenaires (Excel et CSV).

**Provenance** : script écrit pour SortirIA (14/10/2026). Les colonnes viennent
de `TEMPLATE_COLUMNS` (`src/event/partner_mapper.py`) : le modèle et le lecteur
ne peuvent donc pas diverger.

Le classeur contient trois onglets :
- « Programme » : à remplir, une ligne par séance (seul onglet importé) ;
- « Exemple » : deux lignes remplies, clairement fictives ;
- « Aide » : le sens de chaque colonne.

Lancement : uv run python -m scripts.creer_modele_partenaire
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from config.paths import BASE_DIR
from src.event.partner_mapper import TEMPLATE_COLUMNS

OUTPUT_DIR = BASE_DIR / "docs" / "modeles"

# Exemple volontairement fictif (« exemple » dans le nom) : il ne doit jamais
# être pris pour une vraie programmation.
EXAMPLE_ROWS = [
    {
        "titre": "Atelier jeux en famille (exemple)", "date": "21/11/2026", "heure_debut": "14:30",
        "heure_fin": "17:00", "lieu": "Ludothèque Exemple", "adresse": "1 rue de l'Exemple",
        "code_postal": "34200", "commune": "Sète", "prix": "3 €", "gratuit": "non", "public": "famille",
        "age_min": "4", "age_max": "12", "categorie": "Atelier",
        "description": "Jeux de société et de construction pour petits et grands.",
        "lien": "https://exemple.fr/atelier",
    },
    {
        "titre": "Soirée DJ électro (exemple)", "date": "28/11/2026", "heure_debut": "23:00",
        "heure_fin": "05:00", "lieu": "Club Exemple", "commune": "Montpellier", "prix": "10 à 15 €",
        "public": "soirée", "age_min": "18", "categorie": "Soirée DJ",
        "description": "Deux DJ résidents, entrée gratuite avant minuit.", "lien": "https://exemple.fr/soiree",
    },
]
RULES = [
    "Une ligne = une séance. Un spectacle joué trois soirs = trois lignes.",
    "Ne pas changer les noms de colonnes (ligne 1).",
    "Colonnes obligatoires (en orange) : titre, date, lieu, commune.",
    "Chaque séance doit avoir un lien, sauf si le site du partenaire est connu de SortirIA.",
    "Les séances passées sont ignorées : inutile de les effacer.",
    "Pas de photos : seuls les textes sont repris.",
]


def write_excel(path: Path) -> None:
    """Classeur Excel mis en forme, avec listes déroulantes."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = [name for name, _, _ in TEMPLATE_COLUMNS]
    required_fill = PatternFill("solid", fgColor="F4B183")
    optional_fill = PatternFill("solid", fgColor="D9E1F2")
    workbook = Workbook()

    def header_sheet(sheet) -> None:
        sheet.append(headers)
        for cell, (_, required, _) in zip(sheet[1], TEMPLATE_COLUMNS, strict=True):
            cell.font = Font(bold=True)
            cell.fill = required_fill if required else optional_fill
            sheet.column_dimensions[cell.column_letter].width = 32 if cell.value == "description" else 16
        sheet.freeze_panes = "A2"

    programme = workbook.active
    programme.title = "Programme"
    header_sheet(programme)
    choices = {"gratuit": '"oui,non"', "statut": '"complet,annulé,reporté"', "public": '"famille,soirée,tout public,adultes"'}
    for name, formula in choices.items():
        letter = programme.cell(row=1, column=headers.index(name) + 1).column_letter
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.add(f"{letter}2:{letter}1000")
        programme.add_data_validation(validation)

    example = workbook.create_sheet("Exemple")
    header_sheet(example)
    for row in EXAMPLE_ROWS:
        example.append([row.get(name, "") for name in headers])

    help_sheet = workbook.create_sheet("Aide")
    help_sheet.append(["Colonne", "Obligatoire", "À quoi elle sert"])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for name, required, text in TEMPLATE_COLUMNS:
        help_sheet.append([name, "oui" if required else "", text])
    help_sheet.append([])
    help_sheet.append(["Règles"])
    help_sheet.cell(row=help_sheet.max_row, column=1).font = Font(bold=True)
    for rule in RULES:
        help_sheet.append([rule])
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 12
    help_sheet.column_dimensions["C"].width = 80
    for row in help_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical="top")

    workbook.save(path)


def write_csv(path: Path) -> None:
    """Version CSV (séparateur « ; » et BOM : s'ouvre proprement dans Excel)."""

    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle, delimiter=";").writerow([name for name, _, _ in TEMPLATE_COLUMNS])


def main() -> None:
    parser = argparse.ArgumentParser(description="Crée le modèle de fichier partenaire.")
    parser.add_argument("--dossier", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    args.dossier.mkdir(parents=True, exist_ok=True)
    write_excel(args.dossier / "modele_partenaire.xlsx")
    write_csv(args.dossier / "modele_partenaire.csv")
    print(f"Modèle écrit dans {args.dossier} (modele_partenaire.xlsx et .csv).")


if __name__ == "__main__":
    main()
